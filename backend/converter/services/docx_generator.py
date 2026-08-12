import os
from pathlib import Path
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from ..config import OUTPUT_DIR
from ..models.schemas import DocumentStructure, TableData, VisualElement


class DocxGenerator:
    """
    Generates styled, professional Word Documents (.docx) from structured document models.
    Strictly text-and-meaning based: converts every slide into a dedicated Topic ('Slide X: Topic')
    and extracts visual/diagram semantic meaning without embedding raw images.
    """

    def __init__(self, output_dir: Optional[Path] = None):
        self.output_dir = output_dir or OUTPUT_DIR
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, doc_structure: DocumentStructure, filename: str) -> Path:
        doc = Document()

        # Set Standard Page Margins (1 inch all around)
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

            # Footer
            footer = section.footer
            f_p = footer.paragraphs[0]
            f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            f_run = f_p.add_run("AI Presentation-to-Document Converter  |  Validated Content")
            f_run.font.size = Pt(8.5)
            f_run.font.color.rgb = RGBColor(148, 163, 184)

        # 1. Document Title & Subtitle Banner
        self._add_title_banner(doc, doc_structure.title, doc_structure.subtitle)

        # 2. Executive Summary Callout Box
        if doc_structure.executive_summary:
            self._add_executive_summary(doc, doc_structure.executive_summary)

        # 3. Document Sections (Corresponding to Slides: Slide 1: Topic, Slide 2: Topic...)
        for idx, sec in enumerate(doc_structure.sections, start=1):
            self._add_section(doc, sec, idx)

        # Output file path
        safe_filename = "".join(c for c in filename if c.isalnum() or c in ("-", "_", " ")).strip()
        if not safe_filename:
            safe_filename = "converted_document"
        out_path = self.output_dir / f"{safe_filename}.docx"
        doc.save(str(out_path))

        return out_path

    def _add_title_banner(self, doc: Document, title: str, subtitle: Optional[str]):
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(12)
        p_title.paragraph_format.space_after = Pt(4)
        run_title = p_title.add_run(title)
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(26)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)  # Slate 900

        if subtitle:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(18)
            run_sub = p_sub.add_run(subtitle)
            run_sub.font.name = "Calibri"
            run_sub.font.size = Pt(13)
            run_sub.font.italic = True
            run_sub.font.color.rgb = RGBColor(71, 85, 105)  # Slate 600

        # Divider line
        p_hr = doc.add_paragraph()
        p_hr.paragraph_format.space_after = Pt(14)
        p_hr_border = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="12" w:space="1" w:color="2563EB"/></w:pBdr>')
        p_hr._p.get_or_add_pPr().append(p_hr_border)

    def _add_executive_summary(self, doc: Document, summary: str):
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(8)
        p_head.paragraph_format.space_after = Pt(4)
        run_h = p_head.add_run("Executive Summary")
        run_h.font.name = "Calibri"
        run_h.font.size = Pt(15)
        run_h.font.bold = True
        run_h.font.color.rgb = RGBColor(37, 99, 235)  # Blue 600

        # Callout Table for Executive Summary
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        self._set_cell_background(cell, "F8FAFC")
        self._set_cell_left_border(cell, "2563EB", sz="24")

        p_box = cell.paragraphs[0]
        p_box.paragraph_format.space_before = Pt(6)
        p_box.paragraph_format.space_after = Pt(6)
        p_box.paragraph_format.left_indent = Inches(0.15)
        p_box.paragraph_format.right_indent = Inches(0.15)
        run_box = p_box.add_run(summary)
        run_box.font.name = "Calibri"
        run_box.font.size = Pt(10.5)
        run_box.font.color.rgb = RGBColor(51, 65, 85)

        # Space after summary
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(12)

    def _add_section(self, doc: Document, sec, sec_num: int):
        # Section Heading: Always "Slide X: <Topic/Title>"
        clean_title = sec.section_title.strip() if sec.section_title else f"Topic {sec_num}"
        if clean_title.lower().startswith("slide"):
            heading_text = clean_title
        else:
            heading_text = f"Slide {sec.slide_number}: {clean_title}"

        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(18)
        p_sec.paragraph_format.space_after = Pt(6)
        run_sec = p_sec.add_run(heading_text)
        run_sec.font.name = "Calibri"
        run_sec.font.size = Pt(16)
        run_sec.font.bold = True
        run_sec.font.color.rgb = RGBColor(30, 41, 59)  # Slate 800

        # Diagram Callouts if any
        for callout in sec.diagram_callouts:
            self._add_diagram_callout(doc, callout)

        # Paragraphs & Bullets
        for p in sec.paragraphs:
            if p.type == "heading_2":
                p_h2 = doc.add_paragraph()
                p_h2.paragraph_format.space_before = Pt(10)
                p_h2.paragraph_format.space_after = Pt(4)
                r = p_h2.add_run(p.text)
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(71, 85, 105)
            elif p.type == "bullet":
                p_b = doc.add_paragraph(style="List Bullet")
                p_b.paragraph_format.space_before = Pt(2)
                p_b.paragraph_format.space_after = Pt(2)
                if p.level > 0:
                    p_b.paragraph_format.left_indent = Inches(0.25 * (p.level + 1))
                r = p_b.add_run(p.text)
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(51, 65, 85)
            elif p.type == "diagram_analysis" or p.type == "callout_insight":
                self._add_insight_box(doc, p.text)
            else:
                p_body = doc.add_paragraph()
                p_body.paragraph_format.space_before = Pt(4)
                p_body.paragraph_format.space_after = Pt(6)
                p_body.paragraph_format.line_spacing = 1.15
                r = p_body.add_run(p.text)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(30, 41, 59)

        # Visual Elements Semantic Meaning (NO raw images inserted)
        for vis in sec.visual_elements:
            self._add_visual_meaning_box(doc, vis)

        # Tables
        for table_data in sec.tables:
            self._add_table(doc, table_data)

    def _add_diagram_callout(self, doc: Document, text: str):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        self._set_cell_background(cell, "EFF6FF")  # Blue-50
        self._set_cell_left_border(cell, "3B82F6", sz="20")

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.1)

        r_badge = p.add_run("DIAGRAM FLOW: ")
        r_badge.font.bold = True
        r_badge.font.size = Pt(9.5)
        r_badge.font.color.rgb = RGBColor(29, 78, 216)

        r_text = p.add_run(text)
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = RGBColor(30, 58, 138)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def _add_insight_box(self, doc: Document, text: str):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        self._set_cell_background(cell, "F1F5F9")  # Slate-100
        self._set_cell_left_border(cell, "64748B", sz="16")

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.1)

        r = p.add_run(text)
        r.font.italic = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(71, 85, 105)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def _add_visual_meaning_box(self, doc: Document, vis: VisualElement):
        """
        Renders visual meaning and diagrams purely in structured text and callouts.
        Guarantees NO binary image files are embedded.
        """
        has_content = (
            vis.vision_description
            or vis.meaningful_relationships
            or vis.ocr_text
        )
        if not has_content:
            return

        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        self._set_cell_background(cell, "F0FDF4")  # Green/Emerald 50
        self._set_cell_left_border(cell, "059669", sz="20")  # Emerald 600

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.right_indent = Inches(0.1)

        elem_type_label = (vis.type or "visual").replace("_", " ").upper()
        r_badge = p.add_run(f"VISUAL MEANING & ANALYSIS ({elem_type_label}):\n")
        r_badge.font.bold = True
        r_badge.font.size = Pt(9.5)
        r_badge.font.color.rgb = RGBColor(4, 120, 87)

        if vis.vision_description:
            r_desc = p.add_run(vis.vision_description + "\n")
            r_desc.font.size = Pt(10)
            r_desc.font.color.rgb = RGBColor(30, 41, 59)

        if vis.meaningful_relationships:
            p_rel = cell.add_paragraph()
            p_rel.paragraph_format.space_before = Pt(2)
            p_rel.paragraph_format.space_after = Pt(2)
            p_rel.paragraph_format.left_indent = Inches(0.1)
            r_rel_lbl = p_rel.add_run("• Flow & Connections: ")
            r_rel_lbl.font.bold = True
            r_rel_lbl.font.size = Pt(9.5)
            r_rel_lbl.font.color.rgb = RGBColor(4, 120, 87)
            r_rel = p_rel.add_run(vis.meaningful_relationships)
            r_rel.font.size = Pt(9.5)
            r_rel.font.color.rgb = RGBColor(51, 65, 85)

        if vis.ocr_text:
            p_ocr = cell.add_paragraph()
            p_ocr.paragraph_format.space_before = Pt(2)
            p_ocr.paragraph_format.space_after = Pt(4)
            p_ocr.paragraph_format.left_indent = Inches(0.1)
            r_ocr_lbl = p_ocr.add_run("• Extracted Data / Labels: ")
            r_ocr_lbl.font.bold = True
            r_ocr_lbl.font.size = Pt(9.5)
            r_ocr_lbl.font.color.rgb = RGBColor(4, 120, 87)
            r_ocr = p_ocr.add_run(vis.ocr_text)
            r_ocr.font.size = Pt(9.5)
            r_ocr.font.color.rgb = RGBColor(51, 65, 85)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def _add_table(self, doc: Document, table_data: TableData):
        if not table_data.headers and not table_data.rows:
            return

        cols = len(table_data.headers) if table_data.headers else (len(table_data.rows[0]) if table_data.rows else 1)
        rows_count = (1 if table_data.headers else 0) + len(table_data.rows)

        table = doc.add_table(rows=rows_count, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        cur_row = 0
        if table_data.headers:
            for c_idx, head in enumerate(table_data.headers):
                cell = table.cell(cur_row, c_idx)
                self._set_cell_background(cell, "1E293B")  # Dark Slate
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(head)
                r.font.bold = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(255, 255, 255)
            cur_row += 1

        for r_idx, row in enumerate(table_data.rows):
            for c_idx, val in enumerate(row):
                if c_idx < cols:
                    cell = table.cell(cur_row, c_idx)
                    bg = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
                    self._set_cell_background(cell, bg)
                    p = cell.paragraphs[0]
                    r = p.add_run(val)
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(51, 65, 85)
            cur_row += 1

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def _set_cell_background(self, cell, hex_color: str):
        shading = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def _set_cell_left_border(self, cell, hex_color: str, sz: str = "24"):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
